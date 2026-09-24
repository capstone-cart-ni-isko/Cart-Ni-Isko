import docx
doc = docx.Document(r'C:\cartniisko\Cart-Ni-Isko\backend\cartniisko_srs_schema_api.docx')
f = open(r'C:\cartniisko\Cart-Ni-Isko\backend\docx_dump.txt', 'w')
for p in doc.paragraphs:
    f.write(p.text + '\n')
for t in doc.tables:
    for row in t.rows:
        f.write(' | '.join(c.text for c in row.cells) + '\n')
f.close()
print('Done')
